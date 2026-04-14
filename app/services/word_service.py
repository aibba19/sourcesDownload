from pathlib import Path
from urllib.parse import urlparse

from docx import Document

from app.utils.validators import is_supported_youtube_url


class WordService:
    def extract_youtube_links(self, docx_path: Path) -> list[tuple[int, str]]:
        if docx_path.suffix.lower() != ".docx":
            raise ValueError("Il file Word deve avere estensione .docx")

        if not docx_path.exists():
            raise FileNotFoundError(f"File Word non trovato: {docx_path}")

        document = Document(docx_path)
        results: list[tuple[int, str]] = []

        for table in document.tables:
            for row_index, row in enumerate(table.rows, start=1):
                if len(row.cells) < 2:
                    continue

                second_cell = row.cells[1]
                candidates = self._extract_urls_from_cell(second_cell)

                for url in candidates:
                    if is_supported_youtube_url(url):
                        results.append((row_index, url))

        return results

    def _extract_urls_from_cell(self, cell) -> set[str]:
        urls: set[str] = set()

        for paragraph in cell.paragraphs:
            for token in paragraph.text.split():
                clean = token.strip("()[]<>,.;\"'\n\r\t")
                if clean.startswith(("http://", "https://")):
                    urls.add(clean)

            for rel_id in paragraph._p.xpath(".//w:hyperlink/@r:id"):
                rel = paragraph.part.rels.get(rel_id)
                if rel and getattr(rel, "target_ref", None):
                    target = str(rel.target_ref).strip()
                    parsed = urlparse(target)
                    if parsed.scheme in {"http", "https"}:
                        urls.add(target)

        return urls
