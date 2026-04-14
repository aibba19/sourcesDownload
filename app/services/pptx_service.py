from datetime import UTC, datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from app.domain.models import VideoMetadata


def _add_hyperlink(paragraph, text: str, url: str):
    run = paragraph.add_run()
    run.text = text
    run.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
    run.font.underline = True
    run.hyperlink.address = url


class PptxService:
    def build_metadata_pptx(self, metadata: VideoMetadata, output_dir: Path) -> Path:
        pptx_path = output_dir / "fonti.pptx"
        if pptx_path.exists():
            presentation = Presentation(pptx_path)
        else:
            presentation = Presentation()

        slide_layout = presentation.slide_layouts[5]
        slide = presentation.slides.add_slide(slide_layout)

        title_box = slide.shapes.title
        title_box.text = "Video Metadata"

        text_box = slide.shapes.add_textbox(Pt(40), Pt(100), Pt(880), Pt(360))
        text_frame = text_box.text_frame
        text_frame.clear()

        fields = [
            ("Nome file", metadata.output_file_name),
            ("Titolo video", metadata.title),
            ("Canale YouTube", metadata.channel_name),
            ("Data download UTC", datetime.now(UTC).strftime("%Y-%m-%d %H:%M:%S")),
        ]

        for index, (label, value) in enumerate(fields):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = f"{label}: {value}"
            paragraph.font.size = Pt(20)

        video_paragraph = text_frame.add_paragraph()
        video_paragraph.text = "Link video: "
        _add_hyperlink(video_paragraph, metadata.video_url, metadata.video_url)
        video_paragraph.font.size = Pt(20)

        channel_paragraph = text_frame.add_paragraph()
        channel_paragraph.text = "Link canale: "
        if metadata.channel_url:
            _add_hyperlink(channel_paragraph, metadata.channel_url, metadata.channel_url)
        else:
            channel_paragraph.add_run().text = "Non disponibile"
        channel_paragraph.font.size = Pt(20)

        presentation.save(pptx_path)
        return pptx_path
